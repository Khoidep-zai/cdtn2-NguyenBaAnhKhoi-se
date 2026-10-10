#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Script to convert markdown files to professional DOCX documents using python-docx.
Supports: Headings, Metadata, Bullet lists, Numbered lists, Tables, Code blocks, and Inline formatting.
"""

import sys
import re
from pathlib import Path
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.ns import nsdecls, qn

if hasattr(sys.stdout, "reconfigure"):
    sys.stdout.reconfigure(encoding="utf-8")

def set_cell_background(cell, fill_hex):
    """Sets background color of a table cell."""
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    """Sets cell padding in dxa (1 pt = 20 dxa)."""
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(
        f'<w:tcMar {nsdecls("w")}>'
        f'<w:top w:w="{top}" w:type="dxa"/>'
        f'<w:bottom w:w="{bottom}" w:type="dxa"/>'
        f'<w:left w:w="{left}" w:type="dxa"/>'
        f'<w:right w:w="{right}" w:type="dxa"/>'
        f'</w:tcMar>'
    )
    tcPr.append(tcMar)

def add_formatted_text(p, text):
    """Parses basic markdown inline formatting (**bold**, *italic*, `code`) into runs."""
    # Pattern to match bold, code, italic
    tokens = re.split(r'(\*\*.*?\*\*|`.*?`|\*.*?\*)', text)
    for token in tokens:
        if not token:
            continue
        if token.startswith('**') and token.endswith('**') and len(token) >= 4:
            run = p.add_run(token[2:-2])
            run.bold = True
        elif token.startswith('`') and token.endswith('`') and len(token) >= 2:
            run = p.add_run(token[1:-1])
            run.font.name = 'Consolas'
            run.font.size = Pt(9.5)
            run.font.color.rgb = RGBColor(180, 40, 40)
        elif token.startswith('*') and token.endswith('*') and len(token) >= 2:
            run = p.add_run(token[1:-1])
            run.italic = True
        else:
            p.add_run(token)

def convert_md_file(md_path: Path, docx_path: Path):
    print(f"Converting {md_path.name} -> {docx_path.name}...")
    content = md_path.read_text(encoding='utf-8')
    lines = content.splitlines()

    doc = docx.Document()

    # Set standard page margins (1 inch)
    for section in doc.sections:
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)

    # Style defaults
    style_normal = doc.styles['Normal']
    style_normal.font.name = 'Segoe UI'
    style_normal.font.size = Pt(10.5)
    style_normal.font.color.rgb = RGBColor(30, 41, 59) # Slate 800

    i = 0
    total = len(lines)

    in_code_block = False
    code_lines = []

    while i < total:
        line = lines[i].rstrip()

        # Handle Code Block fences
        if line.startswith('```'):
            if in_code_block:
                # End of code block
                code_text = '\n'.join(code_lines)
                p = doc.add_paragraph()
                p.paragraph_format.left_indent = Inches(0.25)
                p.paragraph_format.right_indent = Inches(0.25)
                p.paragraph_format.space_before = Pt(4)
                p.paragraph_format.space_after = Pt(8)
                run = p.add_run(code_text)
                run.font.name = 'Consolas'
                run.font.size = Pt(9)
                run.font.color.rgb = RGBColor(30, 41, 59)

                # Add light border/background to code block paragraph if possible
                pBdr = parse_xml(
                    f'<w:pBdr {nsdecls("w")}>'
                    f'<w:left w:val="single" w:sz="24" w:space="8" w:color="4F46E5"/>'
                    f'</w:pBdr>'
                )
                p._p.get_or_add_pPr().append(pBdr)

                shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="F8FAFC"/>')
                p._p.get_or_add_pPr().append(shd)

                code_lines = []
                in_code_block = False
            else:
                in_code_block = True
                code_lines = []
            i += 1
            continue

        if in_code_block:
            code_lines.append(lines[i])
            i += 1
            continue

        # Handle Tables
        if line.startswith('|') and line.endswith('|'):
            # Collect all table lines
            table_lines = []
            while i < total and lines[i].strip().startswith('|') and lines[i].strip().endswith('|'):
                table_lines.append(lines[i].strip())
                i += 1
            
            if len(table_lines) >= 2:
                # Parse header, separator, data rows
                header_raw = table_lines[0]
                headers = [c.strip() for c in header_raw.strip('|').split('|')]
                data_rows = []
                for t_line in table_lines[2:]: # skip separator line
                    cells = [c.strip() for c in t_line.strip('|').split('|')]
                    # Clean <br> tags
                    cells = [c.replace('<br>', '\n') for c in cells]
                    data_rows.append(cells)

                table = doc.add_table(rows=len(data_rows) + 1, cols=len(headers))
                table.style = 'Table Grid'
                table.alignment = WD_TABLE_ALIGNMENT.CENTER
                table.autofit = True

                # Format Header Row
                hdr_cells = table.rows[0].cells
                for col_idx, h_text in enumerate(headers):
                    if col_idx < len(hdr_cells):
                        cell = hdr_cells[col_idx]
                        set_cell_background(cell, "4F46E5") # Primary indigo
                        set_cell_margins(cell, top=140, bottom=140, left=180, right=180)
                        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
                        p = cell.paragraphs[0]
                        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
                        p.paragraph_format.space_before = Pt(2)
                        p.paragraph_format.space_after = Pt(2)
                        run = p.add_run(h_text)
                        run.bold = True
                        run.font.name = 'Segoe UI'
                        run.font.size = Pt(10)
                        run.font.color.rgb = RGBColor(255, 255, 255)

                # Format Data Rows
                for row_idx, row_data in enumerate(data_rows):
                    row = table.rows[row_idx + 1]
                    bg_color = "F8FAFC" if row_idx % 2 == 1 else "FFFFFF"
                    for col_idx, val in enumerate(row_data):
                        if col_idx < len(row.cells):
                            cell = row.cells[col_idx]
                            set_cell_background(cell, bg_color)
                            set_cell_margins(cell, top=120, bottom=120, left=160, right=160)
                            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
                            p = cell.paragraphs[0]
                            p.paragraph_format.space_before = Pt(1)
                            p.paragraph_format.space_after = Pt(1)
                            p.paragraph_format.line_spacing = 1.15
                            
                            # Add text with formatting (handles bold, code snippets, etc.)
                            add_formatted_text(p, val)
                            for r in p.runs:
                                if not r.bold and not r.italic and r.font.name != 'Consolas':
                                    r.font.name = 'Segoe UI'
                                    r.font.size = Pt(9.5)

                # Add space after table
                spacer = doc.add_paragraph()
                spacer.paragraph_format.space_after = Pt(8)
                continue

        # Handle Headings
        if line.startswith('# '):
            p = doc.add_paragraph(line[2:].strip(), style='Heading 1')
            p.paragraph_format.space_before = Pt(16)
            p.paragraph_format.space_after = Pt(8)
            for r in p.runs:
                r.font.name = 'Segoe UI'
                r.font.size = Pt(18)
                r.bold = True
                r.font.color.rgb = RGBColor(30, 58, 138) # Dark Navy
            i += 1
            continue

        if line.startswith('## '):
            p = doc.add_paragraph(line[3:].strip(), style='Heading 2')
            p.paragraph_format.space_before = Pt(14)
            p.paragraph_format.space_after = Pt(6)
            for r in p.runs:
                r.font.name = 'Segoe UI'
                r.font.size = Pt(14)
                r.bold = True
                r.font.color.rgb = RGBColor(79, 70, 229) # Primary Indigo
            i += 1
            continue

        if line.startswith('### '):
            p = doc.add_paragraph(line[4:].strip(), style='Heading 3')
            p.paragraph_format.space_before = Pt(10)
            p.paragraph_format.space_after = Pt(4)
            for r in p.runs:
                r.font.name = 'Segoe UI'
                r.font.size = Pt(12)
                r.bold = True
                r.font.color.rgb = RGBColor(15, 23, 42) # Slate 900
            i += 1
            continue

        if line.startswith('#### '):
            p = doc.add_paragraph(line[5:].strip(), style='Heading 4')
            p.paragraph_format.space_before = Pt(8)
            p.paragraph_format.space_after = Pt(3)
            for r in p.runs:
                r.font.name = 'Segoe UI'
                r.font.size = Pt(11)
                r.bold = True
                r.font.color.rgb = RGBColor(51, 65, 85)
            i += 1
            continue

        # Handle Horizontal Rules
        if line.strip() in ['---', '***', '___']:
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(6)
            p.paragraph_format.space_after = Pt(6)
            pBdr = parse_xml(
                f'<w:pBdr {nsdecls("w")}>'
                f'<w:bottom w:val="single" w:sz="12" w:space="1" w:color="CBD5E1"/>'
                f'</w:pBdr>'
            )
            p._p.get_or_add_pPr().append(pBdr)
            i += 1
            continue

        # Handle Bullet points
        if line.startswith('- ') or line.startswith('* '):
            item_text = line[2:].strip()
            p = doc.add_paragraph(style='List Bullet')
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.line_spacing = 1.15
            add_formatted_text(p, item_text)
            for r in p.runs:
                if not r.bold and not r.italic and r.font.name != 'Consolas':
                    r.font.name = 'Segoe UI'
                    r.font.size = Pt(10)
            i += 1
            continue

        # Handle Numbered lists
        m_num = re.match(r'^(\d+)\.\s+(.*)', line)
        if m_num:
            item_text = m_num.group(2).strip()
            p = doc.add_paragraph(f"{m_num.group(1)}. ", style='Normal')
            p.paragraph_format.left_indent = Inches(0.25)
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.line_spacing = 1.15
            add_formatted_text(p, item_text)
            for r in p.runs:
                if not r.bold and not r.italic and r.font.name != 'Consolas':
                    r.font.name = 'Segoe UI'
                    r.font.size = Pt(10)
            i += 1
            continue

        # Handle Blockquotes
        if line.startswith('> '):
            quote_text = line[2:].strip()
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.3)
            p.paragraph_format.space_before = Pt(4)
            p.paragraph_format.space_after = Pt(4)
            pBdr = parse_xml(
                f'<w:pBdr {nsdecls("w")}>'
                f'<w:left w:val="single" w:sz="24" w:space="8" w:color="F59E0B"/>'
                f'</w:pBdr>'
            )
            p._p.get_or_add_pPr().append(pBdr)
            shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="FEF3C7"/>')
            p._p.get_or_add_pPr().append(shd)
            add_formatted_text(p, quote_text)
            for r in p.runs:
                r.italic = True
                r.font.size = Pt(10)
                r.font.color.rgb = RGBColor(146, 64, 14)
            i += 1
            continue

        # Handle Normal Paragraphs
        if line.strip():
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(4)
            p.paragraph_format.line_spacing = 1.2
            add_formatted_text(p, line)
            for r in p.runs:
                if not r.bold and not r.italic and r.font.name != 'Consolas':
                    r.font.name = 'Segoe UI'
                    r.font.size = Pt(10.5)

        i += 1

    doc.save(docx_path)
    print(f"✅ Successfully generated {docx_path.name} ({len(doc.paragraphs)} paras, {len(doc.tables)} tables)")

def main():
    docs_dir = Path(__file__).parent
    files_to_sync = [
        ("srs.md", "srs.docx"),
        ("api-contract.md", "api-contract.docx"),
        ("architecture.md", "architecture.docx"),
        ("deployment-guide.md", "deployment-guide.docx"),
        ("phieu-pham-vi.md", "phieu-pham-vi.docx"),
        ("ai-disclosure.md", "ai-disclosure.docx"),
    ]

    for md_name, docx_name in files_to_sync:
        md_file = docs_dir / md_name
        docx_file = docs_dir / docx_name
        if md_file.exists():
            convert_md_file(md_file, docx_file)
        else:
            print(f"⚠️ Skipped {md_name} (file not found)")

if __name__ == '__main__':
    main()
